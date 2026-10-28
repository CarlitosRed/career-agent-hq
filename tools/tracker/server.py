# Career Agent HQ tracker: local server. Signed: Claudio · Career PM (Claude), 2026-10-28.
# Serves index.html and reads/writes work/applications-tracker.xlsx (the xlsx stays the source of truth,
# so the daily rerun and this page edit the same file). Local only (127.0.0.1). Never applies anywhere.
#
# Env vars (optional, for testing against a copy):
#   TRACKER_XLSX  path to the workbook   (default: ../../work/applications-tracker.xlsx)
#   TRACKER_PORT  port to listen on      (default: 8765)
#
# Writes are safe: the workbook is saved to a temp file next to it, re-opened to verify, then swapped in
# with os.replace. One backup per day is kept in .tracker-backups/ beside the workbook (last 7 days).
import json, os, re, shutil, sys, threading, webbrowser
from copy import copy
from datetime import date
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

import openpyxl

HERE = Path(__file__).resolve().parent
XLSX = Path(os.environ.get("TRACKER_XLSX") or HERE.parent.parent / "work" / "applications-tracker.xlsx").resolve()
PORT = int(os.environ.get("TRACKER_PORT") or 8765)
BACKUPS = XLSX.parent / ".tracker-backups"
LOCK = threading.Lock()
# Same order as the data-validation list in the workbook.
STATUSES = ("Suggested (not applied)", "Message drafted", "CV ready", "Applied", "Interview", "Rejected", "Offer", "Skipped")
ADD_FIELDS = ("Company", "Role", "Where", "Posted", "Score", "Verdict", "My status", "Date applied", "CV used", "Source", "Link", "Notes")


class Conflict(Exception):
    pass


def cell(v):
    if v is None:
        return ""
    return v.isoformat()[:10] if hasattr(v, "isoformat") else str(v)


def header(ws):
    return [c.value for c in ws[1]]


def load():
    wb = openpyxl.load_workbook(XLSX)
    jobs, contacts, sheets = [], [], []
    for ws in wb:
        head = header(ws)
        if "My status" in head:
            check_col = next((str(h) for h in head if h and str(h).startswith("Open check")), "")
            m = re.search(r"\d{4}-\d{2}-\d{2}", check_col)
            sheets.append({"name": ws.title, "columns": [str(h) for h in head if h]})
            for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if len(row) < 2 or not row[1]:
                    continue
                d = {str(h): cell(v) for h, v in zip(head, row) if h}
                d["_sheet"], d["_row"] = ws.title, n
                d["_check"] = d.get(check_col, "") if check_col else ""
                d["_checkdate"] = m.group(0) if m else ""
                jobs.append(d)
        elif ws.title == "Contacts":
            for n, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
                if row and row[0]:
                    d = {str(h): cell(v) for h, v in zip(head, row) if h}
                    d["_row"] = n
                    contacts.append(d)
    return {"jobs": jobs, "contacts": contacts, "sheets": sheets, "statuses": STATUSES,
            "today": date.today().isoformat(), "version": version(), "file": XLSX.name}


_CACHE = {"version": None, "body": ""}


def cached():
    """The workbook is parsed once per file version (openpyxl is slow); 'today' is refreshed every call."""
    v = version()
    if v != _CACHE["version"] or not _CACHE["body"]:
        _CACHE["body"], _CACHE["version"] = load(), v
    d = dict(_CACHE["body"], today=date.today().isoformat(), version=v)
    return json.dumps(d)


def version():
    try:
        return str(XLSX.stat().st_mtime_ns)
    except OSError:
        return ""


def excel_open():
    return (XLSX.parent / ("~$" + XLSX.name)).exists()


def backup():
    BACKUPS.mkdir(exist_ok=True)
    today = BACKUPS / f"{XLSX.stem}-{date.today().isoformat()}{XLSX.suffix}"
    if not today.exists():
        shutil.copy2(XLSX, today)
    olds = sorted(BACKUPS.glob(f"{XLSX.stem}-*{XLSX.suffix}"))
    for old in olds[:-7]:
        old.unlink(missing_ok=True)


def save(wb):
    """Save to a temp file, verify it opens with the same sheets, then atomically replace the original."""
    if excel_open():
        raise PermissionError("open in Excel")
    tmp = XLSX.with_name(f".{XLSX.stem}.saving-{os.getpid()}{XLSX.suffix}")
    try:
        wb.save(tmp)
        check = openpyxl.load_workbook(tmp, read_only=True)
        same = check.sheetnames == wb.sheetnames
        check.close()
        if not same:
            raise RuntimeError("verification failed; workbook left unchanged")
        backup()
        os.replace(tmp, XLSX)
    finally:
        if tmp.exists():
            tmp.unlink(missing_ok=True)


def expect_row(ws, head, row, p):
    """Guard against the daily rerun having reordered rows since the page loaded."""
    exp = p.get("expect") or {}
    for k, v in exp.items():
        if k in head and cell(ws.cell(row, head.index(k) + 1).value) != str(v):
            raise Conflict("The spreadsheet changed since this page loaded. Refresh and try again.")


def update(p):
    sheet, row = p["sheet"], int(p["row"])
    wb = openpyxl.load_workbook(XLSX)
    if sheet not in wb.sheetnames:
        raise ValueError("unknown sheet")
    ws = wb[sheet]
    head = header(ws)
    if row < 2 or row > ws.max_row:
        raise ValueError("bad row")
    expect_row(ws, head, row, p)
    if sheet == "Contacts":
        col = {"status": "Status", "date": "Date sent", "note": "Notes"}
    else:
        if "My status" not in head:
            raise ValueError("not a job sheet")
        col = {"status": "My status", "date": "Date applied", "note": "Notes"}
        if "status" in p and p["status"] not in STATUSES:
            raise ValueError("bad status")
    if "date" in p and p["date"] and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", str(p["date"])):
        raise ValueError("date must be YYYY-MM-DD")
    for key in ("status", "date", "note"):
        if key in p:
            if col[key] not in head:
                raise ValueError(f"column {col[key]} missing")
            val = p[key]
            if key == "date" and val:
                val = date.fromisoformat(val)
            # assign .value directly: ws.cell(r, c, None) silently ignores None, so clearing would not work
            ws.cell(row, head.index(col[key]) + 1).value = val if val not in ("", None) else None
    save(wb)


def add(p):
    sheet = p["sheet"]
    wb = openpyxl.load_workbook(XLSX)
    if sheet not in wb.sheetnames:
        raise ValueError("unknown sheet")
    ws = wb[sheet]
    head = header(ws)
    if "My status" not in head:
        raise ValueError("not a job sheet")
    f = {k: str(p.get(k) or "").strip() for k in ADD_FIELDS}
    if not f["Company"] or not f["Role"]:
        raise ValueError("Company and role are required")
    f["My status"] = f["My status"] or STATUSES[0]
    if f["My status"] not in STATUSES:
        raise ValueError("bad status")
    if f["Link"] and not re.match(r"https?://", f["Link"]):
        raise ValueError("Link must start with http:// or https://")
    ci = head.index("Company") + 1
    last = 1
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, ci).value not in (None, ""):
            last = r
    new = last + 1
    # next "#" id, keeping the sheet's prefix (1, 2.. or N1, N2.. or E1, E2..)
    if "#" in head:
        hi, prefix, numeric = 0, "", True
        for r in range(2, last + 1):
            m = re.fullmatch(r"(\D*)(\d+)", str(ws.cell(r, head.index("#") + 1).value or ""))
            if m:
                prefix, numeric = m.group(1), isinstance(ws.cell(r, head.index("#") + 1).value, int)
                hi = max(hi, int(m.group(2)))
        f["#"] = hi + 1 if numeric and not prefix else f"{prefix}{hi + 1}"
    for c, h in enumerate(head, start=1):
        if not h:
            continue
        dst = ws.cell(new, c)
        if last >= 2:
            src = ws.cell(last, c)
            if src.has_style:
                dst._style = copy(src._style)
        v = f.get(str(h), "")
        if h == "Date applied" and v:
            v = date.fromisoformat(v)
        if v not in ("", None):
            dst.value = v
    if last >= 2 and ws.row_dimensions[last].height:
        ws.row_dimensions[new].height = ws.row_dimensions[last].height
    sc = openpyxl.utils.get_column_letter(head.index("My status") + 1)
    for dv in ws.data_validations.dataValidation:
        if any(rng.min_col == rng.max_col == head.index("My status") + 1 for rng in dv.sqref.ranges):
            dv.sqref.add(f"{sc}{new}")
    save(wb)
    saved = {str(h): cell(ws.cell(new, c).value) for c, h in enumerate(head, start=1) if h}
    return new, saved


class H(BaseHTTPRequestHandler):
    def log_message(self, *a):
        pass

    def send(self, code, body, ctype):
        b = body if isinstance(body, bytes) else body.encode("utf-8")
        self.send_response(code)
        self.send_header("Content-Type", ctype)
        self.send_header("Content-Length", str(len(b)))
        self.send_header("Cache-Control", "no-store")
        self.end_headers()
        self.wfile.write(b)

    def local(self):
        if self.headers.get("Host", "").split(":")[0] not in ("127.0.0.1", "localhost"):
            return False
        origin = self.headers.get("Origin")
        return not origin or re.fullmatch(r"https?://(127\.0\.0\.1|localhost)(:\d+)?", origin) is not None

    def do_GET(self):
        if not self.local():
            return self.send(404, "not found", "text/plain")
        path = self.path.split("?")[0]
        try:
            if path == "/api/data":
                with LOCK:
                    self.send(200, cached(), "application/json")
            elif path == "/api/version":
                self.send(200, json.dumps({"version": version()}), "application/json")
            elif path in ("/", "/index.html"):
                self.send(200, (HERE / "index.html").read_bytes(), "text/html; charset=utf-8")
            else:
                self.send(404, "not found", "text/plain")
        except PermissionError:
            self.send(409, json.dumps({"error": f"Can't read {XLSX.name} right now. Try again in a moment."}), "application/json")
        except Exception as e:
            self.send(500, json.dumps({"error": str(e)}), "application/json")

    def do_POST(self):
        path = self.path.split("?")[0]
        if path not in ("/api/update", "/api/add") or not self.local():
            return self.send(404, "not found", "text/plain")
        try:
            p = json.loads(self.rfile.read(int(self.headers.get("Content-Length", 0))))
            with LOCK:
                if path == "/api/update":
                    update(p)
                    out = {"ok": True}
                else:
                    row, saved = add(p)
                    out = {"ok": True, "row": row, "saved": saved}
                out["version"] = version()
            self.send(200, json.dumps(out), "application/json")
        except PermissionError:
            self.send(409, json.dumps({"error": f"Close {XLSX.name} in Excel and try again"}), "application/json")
        except Conflict as e:
            self.send(409, json.dumps({"error": str(e)}), "application/json")
        except Exception as e:
            self.send(400, json.dumps({"error": str(e)}), "application/json")


if __name__ == "__main__":
    srv = ThreadingHTTPServer(("127.0.0.1", PORT), H)
    try:
        with LOCK:
            cached()  # warm the cache so the first page load is fast
    except Exception as e:
        print("Warning: could not read the workbook yet:", e)
    print(f"Tracker at http://127.0.0.1:{PORT}  ->  {XLSX}  (Ctrl+C to stop)")
    if "--no-browser" not in sys.argv:
        webbrowser.open(f"http://127.0.0.1:{PORT}")
    srv.serve_forever()
